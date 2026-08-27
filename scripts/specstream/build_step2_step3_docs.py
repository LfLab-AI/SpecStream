#!/usr/bin/env python3
"""Build the two Chinese Step 2/Step 3 operator documents from Markdown."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "243447"
MUTED = "667085"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", start),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def new_numbering_instance(doc):
    style = doc.styles["List Number"]
    style_num_pr = style._element.pPr.numPr
    base_num_id = int(style_num_pr.numId.val)
    numbering = doc.part.numbering_part.element
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_num_id = int(base_num.find(qn("w:abstractNumId")).get(qn("w:val")))
    next_num_id = (
        max(int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))) + 1
    )
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return next_num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = 0
    num = num_pr.get_or_add_numId()
    num.val = num_id


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in [style.name for style in doc.styles]:
        code = doc.styles.add_style("Code Block", 1)
    else:
        code = doc.styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string("1F2937")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.line_spacing = 1.05


def configure_page(doc, running_label):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(running_label)
    set_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    label = p.add_run("SpecStream 研究记录  |  第 ")
    set_font(label, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    tail = p.add_run(" 页")
    set_font(tail, size=8.5, color=MUTED)


def add_title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(title)
    set_font(run, size=24, color="17365D", bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(subtitle)
    set_font(run, size=12.5, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    values = (("用途", "实现与复现实验"), ("语言", "易懂中文"), ("版本", "2026-08-22"))
    for cell, (label, value) in zip(table.rows[0].cells, values):
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{label}\n")
        set_font(r, size=8, color=MUTED, bold=True)
        r = p.add_run(value)
        set_font(r, size=10.5, color=INK, bold=True)
    set_repeat_header(table.rows[0])
    set_table_geometry(table, [3120, 3120, 3120])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_inline_runs(paragraph, text, code=False):
    pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_font(run, name="Consolas", size=9.5, color="7A3E00")
        elif piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            set_font(run, bold=True)
        else:
            run = paragraph.add_run(piece)
            set_font(
                run, name="Consolas" if code else "Calibri", size=9 if code else 11
            )


def add_code_block(doc, lines):
    for index, line in enumerate(lines):
        p = doc.add_paragraph(style="Code Block")
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = index < len(lines) - 1
        p.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        p.paragraph_format.space_after = Pt(4 if index == len(lines) - 1 else 0)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F4F7")
        p_pr.append(shd)
        add_inline_runs(p, line or " ", code=True)


def add_table(doc, rows):
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        set_row_cant_split(table.rows[row_index])
        for col_index in range(col_count):
            value = values[col_index] if col_index < len(values) else ""
            cell = table.rows[row_index].cells[col_index]
            if row_index == 0:
                set_cell_shading(cell, TABLE_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if row_index == 0:
                p.paragraph_format.keep_with_next = True
            add_inline_runs(p, value)
            for run in p.runs:
                run.font.size = Pt(9.2 if col_count >= 6 else 9.8)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    set_repeat_header(table.rows[0])
    if col_count == 2:
        widths = [1700, TABLE_WIDTH_DXA - 1700]
    elif col_count == 4:
        widths = [1050, 2350, 1600, 4360]
    elif col_count == 8:
        widths = [860, 850, 1120, 840, 1150, 1350, 1200, 1990]
    else:
        base = TABLE_WIDTH_DXA // col_count
        widths = [base] * col_count
        widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def parse_markdown(doc, path):
    lines = path.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    subtitle = (
        "实现边界、文件分工、运行开关与验证方法"
        if "修改说明" in title
        else "单卡与多卡实验矩阵、命令、指标和通过条件"
    )
    add_title_block(doc, title, subtitle)

    index = 1
    in_code = False
    code_lines = []
    number_num_id = None
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            number_num_id = None
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            number_num_id = None
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = []
            for row_index, raw in enumerate(table_lines):
                values = [item.strip() for item in raw.strip("|").split("|")]
                if row_index == 1 and all(
                    re.fullmatch(r":?-{3,}:?", item) for item in values
                ):
                    continue
                rows.append(values)
            add_table(doc, rows)
            continue
        if line.startswith("### "):
            number_num_id = None
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, line[4:])
        elif line.startswith("## "):
            number_num_id = None
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, line[3:])
        elif re.match(r"^\d+\. ", line):
            if number_num_id is None:
                number_num_id = new_numbering_instance(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, number_num_id)
            add_inline_runs(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            number_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:])
        elif line.strip():
            number_num_id = None
            p = doc.add_paragraph()
            add_inline_runs(p, line.strip())
        index += 1


def build(markdown_path, output_path, running_label):
    doc = Document()
    configure_styles(doc)
    configure_page(doc, running_label)
    parse_markdown(doc, markdown_path)
    doc.core_properties.title = markdown_path.stem
    doc.core_properties.subject = "SpecStream Step 2 and Step 3"
    doc.core_properties.author = "SpecStream project"
    doc.core_properties.comments = (
        "Generated from the version-controlled Markdown guide."
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    repo = Path(__file__).resolve().parents[2]
    workspace = repo.parent
    docs = repo / "docs" / "specstream_spectre"
    outputs = [
        (
            docs / "STEP2_STEP3_CODE_CHANGE_GUIDE_CN.md",
            workspace / "SpecStream_第二点与第三点代码修改说明_易懂版.docx",
            "SpecStream | 第二点与第三点代码修改说明",
        ),
        (
            docs / "STEP2_STEP3_COMPLETE_TEST_GUIDE_CN.md",
            workspace
            / "SpecStream_第二点与第三点完整测试文档_含分卡Baseline_易懂版.docx",
            "SpecStream | 第二点与第三点完整测试文档",
        ),
    ]
    for source, output, label in outputs:
        build(source, output, label)
        print(output)


if __name__ == "__main__":
    main()
